#!/usr/bin/env python3
"""Build property/data/pilot-schedule.json from the NJ DCA PILOT Database and Viewer 2026.

Source: https://www.nj.gov/dca/dlgs/misc_docs/2026/PILOT%20Database%20and%20Viewer%202026.xlsx
sheet "Raw Data from UFBs" (one row per reported PILOT, 2,876 rows in 290 towns).

Per municipality (keyed by Treasury district code, the first four characters of a PAMS PIN):
  payment_schedule      reported annual PILOT billing still under a dated agreement, 2026-2035,
                        holding each agreement's latest reported billing flat until its end year
  revenue_projection    sum of payment_schedule for 2026-2030 (five years)
  forecast_year         latest reported agreement end year that is 2026 or later
  municipal_share_pct   reported PILOT billing as a percent of the municipal budget
                        (pilot_billing and municipal_budget from exempt-pilot.json)
  dated_billing_share_pct  percent of the town's reported PILOT billing whose agreement
                        has a usable end date; the rest is "end date not reported"

End dates are read from real dates, M/D/YYYY text and bare years (a bare year means Dec 31).
Years before 2000 or after 2100 and text such as "ongoing", "N/A" or "Max 50 yrs." are treated
as not reported. Billing text such as "--" or "No data" counts as no billing.

Usage: python3 property/scripts/build_pilot_schedule.py path/to/PILOT_Database_2026.xlsx [--check]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "property/data/pilot-schedule.json"
CROSSWALK = ROOT / "property/data/nj-district-crosswalk.json"
EXEMPT_PILOT = ROOT / "property/data/exempt-pilot.json"
SOURCE_URL = "https://www.nj.gov/dca/dlgs/misc_docs/2026/PILOT%20Database%20and%20Viewer%202026.xlsx"
FIRST_YEAR = 2026
SCHEDULE_YEARS = list(range(FIRST_YEAR, FIRST_YEAR + 10))
PROJECTION_YEARS = SCHEDULE_YEARS[:5]


def end_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        d = value.date()
    elif isinstance(value, int):
        d = dt.date(value, 12, 31) if 1000 <= value <= 9999 else None
    elif isinstance(value, str):
        text = value.strip()
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
        if m:
            try:
                d = dt.date(int(m[3]), int(m[1]), int(m[2]))
            except ValueError:
                d = None
        elif re.fullmatch(r"\d{4}", text):
            d = dt.date(int(text), 12, 31)
        else:
            d = None
    else:
        d = None
    if d is None or not 2000 <= d.year <= 2100:
        return None
    return d


def billing(value) -> float:
    return float(value) if isinstance(value, (int, float)) and value > 0 else 0.0


def build(xlsx: Path) -> dict:
    crosswalk = json.loads(CROSSWALK.read_text())["districts"]
    by_dca = {v["dca_code"]: k for k, v in crosswalk.items()}
    exempt = json.loads(EXEMPT_PILOT.read_text())["municipalities"]

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Raw Data from UFBs"]
    header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    expected = ["DCA Municode", "Municipality", "County", "Project Name", "Type of Project",
                "Agreement Start Date", "Agreement End Date", "PILOT Billing", "Assessed Value",
                "Taxes if Billed at PY Rate", "Latest UFB Submission Year"]
    if header[: len(expected)] != expected:
        raise SystemExit(f"unexpected header: {header[:len(expected)]}")

    towns: dict[str, dict] = defaultdict(lambda: {"rows": 0, "dated_rows": 0, "billing": 0.0,
                                                  "dated_billing": 0.0, "schedule": defaultdict(float),
                                                  "end_years": []})
    unmatched = set()
    total_rows = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        code = str(row[0] or "").strip()
        if not code:
            continue
        total_rows += 1
        code = code.zfill(4)
        district = by_dca.get(code)
        if district is None:
            unmatched.add(code)
            continue
        t = towns[district]
        amount = billing(row[7])
        end = end_date(row[6])
        t["rows"] += 1
        t["billing"] += amount
        if end is None:
            continue
        t["dated_rows"] += 1
        t["dated_billing"] += amount
        if end.year >= FIRST_YEAR:
            t["end_years"].append(end.year)
        for year in SCHEDULE_YEARS:
            if end.year >= year:
                t["schedule"][year] += amount
    if unmatched:
        raise SystemExit(f"DCA codes missing from the crosswalk: {sorted(unmatched)}")

    out = {}
    for district in sorted(towns):
        t = towns[district]
        ex = exempt.get(district) or {}
        name = crosswalk[district]["name"]
        rec = {
            "name": name,
            "county": crosswalk[district]["county"],
            "reported_rows": t["rows"],
            "dated_rows": t["dated_rows"],
            "reported_billing": round(t["billing"], 2),
            "dated_billing": round(t["dated_billing"], 2),
            "dated_billing_share_pct": round(100 * t["dated_billing"] / t["billing"], 1) if t["billing"] > 0 else None,
            "payment_schedule": None,
            "revenue_projection": None,
            "forecast_year": max(t["end_years"]) if t["end_years"] else None,
            "municipal_share_pct": None,
        }
        if t["dated_rows"]:
            rec["payment_schedule"] = [{"year": y, "value": round(t["schedule"][y], 2)} for y in SCHEDULE_YEARS]
            rec["revenue_projection"] = round(sum(t["schedule"][y] for y in PROJECTION_YEARS), 2)
        pilot_billing, budget = ex.get("pilot_billing"), ex.get("municipal_budget")
        if isinstance(pilot_billing, (int, float)) and isinstance(budget, (int, float)) and budget > 0:
            rec["municipal_share_pct"] = round(100 * pilot_billing / budget, 2)
        out[district] = rec

    return {
        "schema_version": 1,
        "district_key": "modiv_pams",
        "source_year": 2025,
        "release_year": 2026,
        "schedule_years": SCHEDULE_YEARS,
        "projection_years": PROJECTION_YEARS,
        "source": {"agency": "NJ Department of Community Affairs, DLGS",
                   "label": "PILOT Database and Viewer 2026, Raw Data from UFBs", "url": SOURCE_URL},
        "notes": [
            "Schedule values hold each agreement's latest reported PILOT billing flat through its reported end year.",
            "Rows without a usable end date are left out of the schedule and counted as end date not reported.",
            "municipal_share_pct uses pilot_billing and municipal_budget from exempt-pilot.json.",
        ],
        "totals": {"reported_rows": total_rows, "municipalities": len(out),
                   "dated_rows": sum(t["dated_rows"] for t in towns.values())},
        "municipalities": out,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--check", action="store_true", help="fail if the committed file is out of date")
    args = ap.parse_args()
    data = build(args.xlsx)
    text = json.dumps(data, indent=1, sort_keys=False) + "\n"
    if args.check:
        current = OUT.read_text() if OUT.exists() else ""
        if current != text:
            print(f"{OUT.relative_to(ROOT)} is out of date", file=sys.stderr)
            return 1
        return 0
    OUT.write_text(text)
    print(f"wrote {OUT.relative_to(ROOT)}: {data['totals']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
