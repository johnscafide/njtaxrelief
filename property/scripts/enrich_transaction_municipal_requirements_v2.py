#!/usr/bin/env python3
"""Enrich Watchdog municipal CO/CCO + smoke/fire requirements for all 564 NJ municipalities.

This v2 pass fixes three gaps in the original extractor:
1. DCA ordinance links are county-aware so same-named municipalities cannot collide.
2. Ordinance text is actually fetched and parsed instead of being stored as a link only.
3. Coverage metadata distinguishes a found process from extracted municipality-specific checklist detail.

Policy remains conservative: missing, unreadable, or ambiguous official material never means
"not required" and never creates a clearance conclusion.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib.util
import io
import json
import pathlib
import re
from collections import defaultdict
from typing import Any

ROOT = pathlib.Path(__file__).resolve().parents[2]
BASE_PATH = ROOT / "property/scripts/extract_transaction_municipal_requirements.py"
NJ_COUNTIES = {
    "ATLANTIC","BERGEN","BURLINGTON","CAMDEN","CAPE MAY","CUMBERLAND","ESSEX",
    "GLOUCESTER","HUDSON","HUNTERDON","MERCER","MIDDLESEX","MONMOUTH","MORRIS",
    "OCEAN","PASSAIC","SALEM","SOMERSET","SUSSEX","UNION","WARREN",
}


def load_base():
    spec = importlib.util.spec_from_file_location("watchdog_municipal_requirements_base", BASE_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError("Could not load base municipal requirements extractor")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def utcnow() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def county_token(value: str) -> str:
    cleaned = re.sub(r"\bCOUNTY\b", "", str(value or "").upper()).strip(" .,-")
    return cleaned if cleaned in NJ_COUNTIES else ""


def ordinance_links_county_aware(base) -> tuple[dict[tuple[str, str], list[str]], dict[str, list[str]]]:
    """Read DCA's ordinance workbook without collapsing same-name municipalities across counties."""
    try:
        raw, _, _ = base.fetch_bytes(base.DCA_ORDINANCE_XLSX, timeout=45, max_bytes=24_000_000)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
    except Exception:
        return {}, {}

    keyed: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_name: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            vals = [base.clean_text(c.value, 300) for c in row]
            urls: list[str] = []
            for cell in row:
                target = getattr(getattr(cell, "hyperlink", None), "target", None)
                if target and str(target).startswith(("http://", "https://")):
                    urls.append(str(target))
                urls += re.findall(r"https?://[^\s<>\"]+", str(cell.value or ""))
            if not urls:
                continue

            county = next((county_token(v) for v in vals if county_token(v)), "")
            name = next(
                (
                    v for v in vals[:8]
                    if v
                    and not re.fullmatch(r"\d{1,6}", v)
                    and not county_token(v)
                    and len(v) > 2
                    and "municip" not in v.lower()
                    and "county" not in v.lower()
                    and "ordinance" not in v.lower()
                ),
                "",
            )
            if not name:
                continue
            norm = base.norm_name(name)
            for url in urls:
                if county and url not in keyed[(norm, county)]:
                    keyed[(norm, county)].append(url)
                if url not in by_name[norm][county]:
                    by_name[norm][county].append(url)

    # Fallback by name only is safe only when the workbook has exactly one county variant.
    unique_name: dict[str, list[str]] = {}
    for norm, variants in by_name.items():
        nonempty = {k: v for k, v in variants.items() if v}
        if len(nonempty) == 1:
            unique_name[norm] = next(iter(nonempty.values()))
    return dict(keyed), unique_name


def choose_urls_v2(base, muni: dict[str, Any], family: str) -> list[dict[str, Any]]:
    family_key = "certificate_of_occupancy" if family == "resale_cco" else "smoke_co_fire"
    candidates = list((muni.get("candidates") or {}).get(family_key, []) or [])
    candidates.sort(key=lambda x: (-int(x.get("score") or 0), len(str(x.get("url") or ""))))
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in candidates:
        url = str(item.get("url") or "")
        if not url.startswith(("http://", "https://")) or url in seen:
            continue
        seen.add(url)
        out.append(item)
        if len(out) >= 8:
            break
    return out


def build_row_v2(base, muni: dict[str, Any], family: str,
                 ordinance_by_key: dict[tuple[str, str], list[str]],
                 ordinance_unique_name: dict[str, list[str]], generated: str) -> dict[str, Any]:
    code = str(muni.get("municipality_code") or "").strip()
    name = str(muni.get("municipality_label") or code)
    county = str(muni.get("county") or "").upper().strip()
    root = str(muni.get("root_url") or "")
    items = choose_urls_v2(base, muni, family)

    norm = base.norm_name(name)
    ordinance_candidates = ordinance_by_key.get((norm, county), []) or ordinance_unique_name.get(norm, [])
    ordinance_url = ordinance_candidates[0] if ordinance_candidates else None

    source_docs: list[dict[str, str]] = []
    combined: list[str] = []
    fetched_urls: list[str] = []

    for item in items:
        url = str(item.get("url") or "")
        text, final = base.source_text(url)
        if text:
            combined.append(text)
            fetched_urls.append(final)
        source_docs.append({"label": base.clean_text(item.get("label") or final or url, 180), "url": final or url})

    ordinance_text = ""
    ordinance_final = ordinance_url
    if ordinance_url:
        ordinance_text, ordinance_final = base.source_text(ordinance_url)
        if ordinance_text:
            combined.append(ordinance_text)
            fetched_urls.append(ordinance_final)
        source_docs.append({"label": "Local code / ordinance source", "url": ordinance_final or ordinance_url})

    all_text = " ".join(combined)
    local_requirements = base.sentence_candidates(all_text, family)
    requirements = list(local_requirements)
    fees = base.fees_from(all_text)
    application, department = base.classify_urls(items, root)

    if family == "smoke_fire_cert":
        baseline = (
            "NJ fire-safety change-of-occupancy compliance must be verified through the applicable "
            "local enforcing agency or municipal occupancy process; Watchdog does not treat missing "
            "local web material as a waiver."
        )
        if baseline not in requirements:
            requirements.insert(0, baseline)
        source_docs.extend([
            {"label": "NJ DCA Fire Solutions", "url": base.DCA_FIRE_HOME},
            {"label": "NJ DCA Smoke/CO certification application", "url": base.DCA_SMOKE_APP},
        ])

    is_explicit = base.explicit_required(all_text, family)
    local_item_count = len(local_requirements)
    official_material_found = bool(items or ordinance_url or application or department)
    details_extracted = local_item_count > 0

    if is_explicit:
        state = "explicit_required"
    elif official_material_found and (details_extracted or application or department):
        state = "official_process_found"
    elif family == "smoke_fire_cert":
        state = "statewide_baseline"
    else:
        state = "verify"

    seen: set[str] = set()
    sources: list[dict[str, str]] = []
    for src in source_docs:
        url = str(src.get("url") or "")
        if url and url not in seen:
            seen.add(url)
            sources.append(src)
    if root and root not in seen:
        sources.append({"label": f"{name} official website", "url": root})
        seen.add(root)
    sources = sources[:14]

    excerpt = " | ".join(requirements[:5])[:1600]
    facts = {
        "code": code,
        "county": county,
        "family": family,
        "state": state,
        "requirements": requirements[:14],
        "fees": fees,
        "application": application,
        "department": department,
        "ordinance": ordinance_final or ordinance_url,
        "sources": sources,
        "municipality_specific_item_count": local_item_count,
    }
    return {
        "municipality_code": code,
        "municipality_name": name,
        "county": county,
        "requirement_key": family,
        "requirement_state": state,
        "title": "Certificate of Occupancy (CO)" if family == "resale_cco" else "Smoke / CO / fire certification",
        "requirements": requirements[:14],
        "fees": fees,
        "application_url": application,
        "department_url": department,
        "ordinance_url": ordinance_final or ordinance_url,
        "source_urls": sources,
        "source_excerpt": excerpt or None,
        "source_hash": hashlib.sha256(json.dumps(facts, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
        "last_verified_at": generated,
        "metadata": {
            "extractor_version": 2,
            "explicit_requirement_language_observed": is_explicit,
            "candidate_count": len(items),
            "municipality_specific_item_count": local_item_count,
            "municipality_specific_details_extracted": details_extracted,
            "official_material_found": official_material_found,
            "ordinance_text_fetched": bool(ordinance_text),
            "ordinance_county_matched": bool(ordinance_by_key.get((norm, county))),
            "fetched_source_count": len(set(fetched_urls)),
            "never_infer_not_required": True,
            "statewide_smoke_source_applied": family == "smoke_fire_cert",
            "dca_ordinance_directory": base.DCA_ORDINANCE_XLSX,
        },
        "updated_at": generated,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--municipal", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    base = load_base()
    with open(args.municipal, encoding="utf-8") as f:
        doc = json.load(f)
    rows = doc.get("results") or []
    codes = {str(r.get("municipality_code") or "") for r in rows}
    if len(codes) != 564:
        raise SystemExit(f"Expected 564 municipality codes, got {len(codes)}")

    generated = utcnow()
    ordinance_by_key, ordinance_unique = ordinance_links_county_aware(base)
    out = [
        build_row_v2(base, row, family, ordinance_by_key, ordinance_unique, generated)
        for row in rows
        for family in ("resale_cco", "smoke_fire_cert")
    ]

    summary: dict[str, Any] = {
        "municipalities": len(codes),
        "rows": len(out),
        "families": {},
    }
    for family in ("resale_cco", "smoke_fire_cert"):
        fam = [r for r in out if r["requirement_key"] == family]
        summary["families"][family] = {
            "states": {
                state: sum(r["requirement_state"] == state for r in fam)
                for state in ("explicit_required", "official_process_found", "statewide_baseline", "verify")
            },
            "with_municipality_specific_items": sum(
                int((r.get("metadata") or {}).get("municipality_specific_item_count") or 0) > 0 for r in fam
            ),
            "with_ordinance_text": sum(bool((r.get("metadata") or {}).get("ordinance_text_fetched")) for r in fam),
        }

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump({"generated_at": generated, "summary": summary, "rows": out}, f, indent=2, ensure_ascii=False)
    print(json.dumps(summary, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
